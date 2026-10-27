#!/usr/bin/env python3
"""
perfgrab - performance data collection for Android games over adb.

For every selected game:
  1. Reads device info (model, SoC, GPU, RAM, Android version).
  2. Measures 1st and 2nd cold launch time (am start -W).
  3. While the game is played, samples FPS, CPU, GPU, RAM and temperatures.
  4. Saves samples.csv, fps_per_second.csv, summary.json, summary.xlsx,
     chart.png and a tab-separated row for a results spreadsheet.

Standard library only. Optional: openpyxl (xlsx), matplotlib (chart).
"""

__version__ = "1.0.0"
import argparse
import bisect
import csv
import json
import os
import re
import shutil
import shlex
import statistics
import subprocess
import sys
import threading
import time
from datetime import datetime

INT64_MAX = 9223372036854775807

TABLE_COLUMNS = [
    "Tester", "Graphic settings", "Device model", "CPU", "GPU", "RAM", "OS",
    "Duration of the test", "1st launch time, sec", "2nd launch time, sec",
    "Average FPS", "Min FPS", "Max FPS", "Avg frame time", "Avg GPU usage",
    "Avg CPU usage", "Battery temperature, before test",
    "Battery temperature, after 30 minutes", "Link with results", "",
    "Available RAM before launch", "Available RAM during gameplay",
    "RAM the game uses", "", "Temp before the test", "Temp after 3 minutes gameplay",
]

GPU_PATHS = [
    "/sys/class/kgsl/kgsl-3d0/gpu_busy_percentage",   # Qualcomm Adreno
    "/sys/class/kgsl/kgsl-3d0/gpubusy",               # Qualcomm Adreno (busy total)
    "/sys/kernel/gpu/gpu_busy",                       # Exynos / Tensor / some Mali
    "/sys/class/misc/mali0/device/utilization",       # Mali
    "/sys/devices/platform/mali.0/utilization",       # Mali (old)
    "/sys/kernel/ged/hal/gpu_utilization",            # MediaTek
    "/sys/module/ged/parameters/gpu_loading",         # MediaTek (old)
    "/proc/mali/utilization",                         # Mali (some kernels)
]

STANDARD_RAM_GB = [1, 1.5, 2, 3, 4, 6, 8, 10, 12, 16, 18, 20, 24, 32]


# --------------------------------------------------------------------------- adb
def find_adb(user_value):
    """adb from --adb, PATH, or the usual install places on macOS / Windows / Linux."""
    if user_value and user_value != "adb":
        return os.path.expanduser(user_value)
    found = shutil.which("adb")
    if found:
        return found
    home = os.path.expanduser("~")
    candidates = [
        os.path.join(home, "Library/Android/sdk/platform-tools/adb"),      # Android Studio on macOS
        "/opt/homebrew/bin/adb",                                            # Homebrew, Apple Silicon
        "/usr/local/bin/adb",                                               # Homebrew, Intel Mac
        "/opt/homebrew/Caskroom/android-platform-tools/latest/platform-tools/adb",
        os.path.join(os.environ.get("LOCALAPPDATA", ""), "Android/Sdk/platform-tools/adb.exe"),
        os.path.join(home, "Android/Sdk/platform-tools/adb"),               # Linux
    ]
    for c in candidates:
        if os.path.isfile(c):
            return c
    sys.exit("adb not found.\n  macOS:   brew install --cask android-platform-tools\n"
             "  or pass the path: --adb ~/Library/Android/sdk/platform-tools/adb")


class Adb:
    def __init__(self, adb_cmd="adb", serial=None):
        self.base = [adb_cmd]
        if serial:
            self.base += ["-s", serial]

    def run(self, *args, timeout=30):
        try:
            p = subprocess.run(self.base + list(args), capture_output=True, timeout=timeout)
        except subprocess.TimeoutExpired:
            return ""
        except FileNotFoundError:
            sys.exit(f"adb not found at {self.base[0]}. Install platform-tools or pass --adb <path>")
        return p.stdout.decode("utf-8", "replace")

    def sh(self, cmd, timeout=30):
        return self.run("shell", cmd, timeout=timeout)

    def prop(self, name):
        return self.sh(f"getprop {name}").strip()


def list_devices(adb_cmd):
    out = Adb(adb_cmd).run("devices")
    devs = []
    for line in out.splitlines()[1:]:
        parts = line.split("\t")
        if len(parts) == 2 and parts[1].strip() == "device":
            devs.append(parts[0].strip())
    return devs


# ------------------------------------------------------------------- device info
def device_info(adb):
    man = adb.prop("ro.product.manufacturer")
    model = adb.prop("ro.product.model")
    release = adb.prop("ro.build.version.release")
    soc = " ".join(x for x in (adb.prop("ro.soc.manufacturer"), adb.prop("ro.soc.model")) if x)
    if not soc:
        soc = (adb.prop("ro.hardware.chipname") or adb.prop("ro.board.platform")
               or adb.prop("ro.hardware"))
    gpu = ""
    gl = adb.sh("dumpsys SurfaceFlinger | grep -i -m1 GLES")
    m = re.search(r"GLES:\s*(.*)", gl)
    if m:
        parts = [x.strip() for x in m.group(1).split(",")]
        idx = next((i for i, x in enumerate(parts) if x.startswith("OpenGL")), len(parts))
        gpu = parts[idx - 1] if idx > 0 else m.group(1).strip()
    mem_total_kb = 0
    m = re.search(r"MemTotal:\s*(\d+)", adb.sh("cat /proc/meminfo"))
    if m:
        mem_total_kb = int(m.group(1))
    gb = mem_total_kb / 1024 / 1024
    ram_gb = next((s for s in STANDARD_RAM_GB if s >= gb), round(gb)) if gb else None
    cores = 0
    present = adb.sh("cat /sys/devices/system/cpu/present").strip()
    for part in present.split(","):
        if "-" in part:
            a, b = part.split("-")
            cores += int(b) - int(a) + 1
        elif part.isdigit():
            cores += 1
    return {
        "manufacturer": man, "model": model, "device": f"{man} {model}".strip(),
        "cpu": soc, "gpu": gpu, "ram_gb": ram_gb, "ram_total_mb": round(mem_total_kb / 1024),
        "os": f"Android {release}", "cores": cores or None,
    }


# ------------------------------------------------------------------------ launch
def launch_component(adb, pkg):
    out = adb.sh(f"cmd package resolve-activity --brief -c android.intent.category.LAUNCHER {pkg}")
    for line in reversed(out.strip().splitlines()):
        line = line.strip()
        if re.match(r"^[\w.]+/[\w.$]+$", line):
            return line
    out = adb.sh(f"dumpsys package {pkg}")
    m = re.search(r"android\.intent\.action\.MAIN:\s*\n\s*\S+\s+([\w.]+/[\w.$]+)", out)
    return m.group(1) if m else None


def cold_start(adb, pkg, component):
    adb.sh(f"am force-stop {pkg}")
    time.sleep(2)
    out = adb.sh(f"am start -W -n {shlex.quote(component)}", timeout=180)
    m = re.search(r"TotalTime:\s*(\d+)", out) or re.search(r"WaitTime:\s*(\d+)", out)
    return round(int(m.group(1)) / 1000, 2) if m else None


def pidof(adb, pkg):
    out = adb.sh(f"pidof {pkg}").split()
    return out[0] if out and out[0].isdigit() else None


# --------------------------------------------------------------------------- FPS
def find_layers(adb, pkg):
    out = adb.sh("dumpsys SurfaceFlinger --list")
    names = {line.strip() for line in out.splitlines() if pkg in line}

    def score(name):
        s = 0
        if "SurfaceView" in name:
            s += 4
        if "BLAST" in name:
            s += 2
        if "Background" in name:
            s -= 10
        if name.startswith(("ActivityRecord", "Task=", "WindowToken")):
            s -= 3
        return -s

    return sorted(names, key=score)


def read_latency(adb, layer):
    out = adb.sh("dumpsys SurfaceFlinger --latency " + shlex.quote(layer), timeout=10)
    lines = out.strip().splitlines()
    if len(lines) < 2 or not lines[0].strip().isdigit():
        return []
    ts = []
    for line in lines[1:]:
        p = line.split()
        if len(p) < 3:
            continue
        try:
            desired, actual, ready = (int(x) for x in p[:3])
        except ValueError:
            continue
        for t in (actual, ready, desired):
            if 0 < t < INT64_MAX:
                ts.append(t)
                break
    return sorted(ts)


FPS_METHOD_NAMES = {
    "latency": "SurfaceFlinger --latency (frame timestamps)",
    "timestats": "SurfaceFlinger --timestats (frames per layer)",
    "sf_counter": "SurfaceFlinger frame counter (whole screen)",
    "gfxinfo": "gfxinfo (HWUI frames)",
}


def timestats_frames(adb, pkg):
    """{layerName: totalFrames} for layers of the game from SurfaceFlinger TimeStats."""
    out = adb.sh("dumpsys SurfaceFlinger --timestats -dump", timeout=15)
    res, layer, package = {}, None, ""
    for line in out.splitlines():
        line = line.strip()
        key, _, val = line.partition("=")
        key, val = key.strip(), val.strip()
        if key == "layerName":
            layer, package = val, ""
        elif key == "packageName":
            package = val
        elif key == "totalFrames" and layer:
            if (pkg in layer or package == pkg) and val.isdigit():
                res[layer] = int(val)
            layer = None
    return res


def sf_frame_counter(adb):
    """Global count of frames composed by SurfaceFlinger (transaction 1013)."""
    out = adb.sh("service call SurfaceFlinger 1013")
    m = re.search(r"Parcel\(\s*([0-9a-fA-F]{8})", out)
    return int(m.group(1), 16) if m else None


def gfxinfo_frames(adb, pkg):
    m = re.search(r"Total frames rendered:\s*(\d+)", adb.sh(f"dumpsys gfxinfo {pkg}", timeout=15))
    return int(m.group(1)) if m else None


class FrameCollector(threading.Thread):
    """
    Collects game frames with the first method that works on this phone:
      1. latency    - per-frame timestamps of the game layer (most precise, 1% low, frame time)
      2. timestats  - frame counter of the game layer
      3. sf_counter - frames composed on the whole screen (fine for a fullscreen game)
      4. gfxinfo    - HWUI frames (works only for non-engine apps)
    """

    def __init__(self, adb, pkg, interval=0.4, log=None, force=None):
        super().__init__(daemon=True)
        self.adb, self.pkg, self.interval = adb, pkg, interval
        self.log = log or (lambda m: None)
        self.force = force
        self.ts = []            # latency mode: device timestamps, ns
        self.counter = []       # counter modes: (wall time, cumulative frames)
        self.total = 0
        self.prev_raw = None
        self.ts_layer = None    # timestats layer
        self.last = 0
        self.layer = None
        self.mode = None
        self.lock = threading.Lock()
        self.stop_ev = threading.Event()

    # ---- latency
    def _add(self, ts, baseline=False):
        with self.lock:
            new = [t for t in ts if t > self.last]
            if new:
                self.last = new[-1]
                if not baseline:
                    self.ts.extend(new)

    def _try_latency(self):
        for cand in find_layers(self.adb, self.pkg):
            ts = read_latency(self.adb, cand)
            if len(ts) >= 2:
                self.layer = cand
                self._add(ts, baseline=(self.last == 0))
                return True
        return False

    # ---- counters
    def _read_counter(self, mode):
        if mode == "timestats":
            d = timestats_frames(self.adb, self.pkg)
            if self.ts_layer:
                return d.get(self.ts_layer)
            return d
        if mode == "sf_counter":
            return sf_frame_counter(self.adb)
        if mode == "gfxinfo":
            return gfxinfo_frames(self.adb, self.pkg)
        return None

    def _try_counter(self, mode):
        if mode == "timestats":
            self.adb.sh("dumpsys SurfaceFlinger --timestats -enable")
        a = self._read_counter(mode)
        if self.stop_ev.wait(1.2):
            return False
        b = self._read_counter(mode)
        if mode == "timestats":
            if not isinstance(a, dict) or not isinstance(b, dict) or not b:
                return False
            deltas = {k: b[k] - a.get(k, 0) for k in b}
            best = max(deltas, key=deltas.get)
            if deltas[best] <= 0:
                return False
            self.ts_layer, b = best, b[best]
        elif a is None or b is None or b <= a:
            return False
        self.prev_raw = b
        with self.lock:
            self.counter.append((time.time(), self.total))
        return True

    def _poll_counter(self):
        v = self._read_counter(self.mode)
        if v is None or isinstance(v, dict):
            return False
        d = v - self.prev_raw
        if d < 0:          # counter reset (layer recreated, timestats cleared)
            d = 0
        self.prev_raw = v
        with self.lock:
            self.total += d
            self.counter.append((time.time(), self.total))
        return True

    def _set_mode(self, mode):
        self.mode = mode
        extra = f" | layer: {self.layer or self.ts_layer}" if (self.layer or self.ts_layer) else ""
        self.log(f"  FPS source: {FPS_METHOD_NAMES[mode]}{extra}")

    def run(self):
        start = time.time()
        misses = 0
        order = [self.force] if self.force else ["latency", "timestats", "sf_counter", "gfxinfo"]
        warned = False
        while not self.stop_ev.is_set():
            if self.mode is None:
                for m in order:
                    if m == "latency":
                        ok = self._try_latency()
                    else:
                        if m != self.force and time.time() - start < 4:
                            break      # give the latency method a few seconds first
                        ok = self._try_counter(m)
                    if ok:
                        self._set_mode(m)
                        break
                if self.mode is None:
                    if not warned and time.time() - start > 15:
                        self.log("  ! FPS: no method works on this device yet. Run FPS diagnostics for details.")
                        warned = True
                    self.stop_ev.wait(2)
                    continue
            elif self.mode == "latency":
                ts = read_latency(self.adb, self.layer)
                if ts:
                    misses = 0
                    self._add(ts)
                else:
                    misses += 1
                    if misses >= 6:
                        self.layer, self.mode, misses = None, None, 0
            else:
                if not self._poll_counter():
                    misses += 1
                    if misses >= 5:
                        self.mode, self.ts_layer, misses = None, None, 0
                else:
                    misses = 0
            self.stop_ev.wait(self.interval if self.mode in (None, "latency") else 1.0)

    def count(self):
        with self.lock:
            return len(self.ts) if self.ts else self.total

    def fps_recent(self, window_sec):
        with self.lock:
            if self.mode == "latency" and len(self.ts) >= 2:
                end = self.ts[-1]
                i = bisect.bisect_right(self.ts, end - int(window_sec * 1e9))
                return (len(self.ts) - i) / window_sec
            if self.mode and self.mode != "latency" and len(self.counter) >= 2:
                t2, f2 = self.counter[-1]
                j = len(self.counter) - 2
                while j > 0 and t2 - self.counter[j][0] < window_sec:
                    j -= 1
                t1, f1 = self.counter[j]
                return (f2 - f1) / (t2 - t1) if t2 > t1 else None
            return None

    def snapshot(self):
        with self.lock:
            return list(self.ts)

    def stats(self):
        with self.lock:
            ts, counter = list(self.ts), list(self.counter)
        if len(ts) >= 3:
            res = fps_stats(ts)
            res["method"] = "latency"
            return res
        if len(counter) >= 3:
            res = counter_stats(counter)
            res["method"] = self.mode or "counter"
            return res
        return {}


def counter_stats(counter):
    (t0, f0), (t1, f1) = counter[0], counter[-1]
    dur = t1 - t0
    if dur <= 0:
        return {}
    per = []
    for (a, fa), (b, fb) in zip(counter, counter[1:]):
        if b > a:
            per.append((fb - fa) / (b - a))
    avg = (f1 - f0) / dur
    return {
        "frames": f1 - f0, "avg_fps": avg,
        "min_fps": round(min(per)) if per else None, "max_fps": round(max(per)) if per else None,
        "avg_frame_time_ms": 1000 / avg if avg else None, "p99_frame_time_ms": None,
        "one_percent_low_fps": None, "frames_over_50ms": None,
        "per_second": [round(x, 1) for x in per],
    }


def fps_diagnostics(adb, pkg):
    """Raw output of every FPS method, for devices where FPS cannot be measured."""
    rep = []
    add = rep.append
    add(f"package: {pkg}")
    add(f"android: {adb.prop('ro.build.version.release')} (sdk {adb.prop('ro.build.version.sdk')}), "
        f"model: {adb.prop('ro.product.manufacturer')} {adb.prop('ro.product.model')}")
    add(f"pidof: {adb.sh(f'pidof {pkg}').strip() or '(not running!)'}")
    add("\n--- SurfaceFlinger --list (lines with package) ---")
    layers = find_layers(adb, pkg)
    add("\n".join(layers) or "(none)")
    for cand in layers[:6]:
        out = adb.sh("dumpsys SurfaceFlinger --latency " + shlex.quote(cand), timeout=10)
        lines = out.strip().splitlines()
        add(f"\n--- --latency '{cand}' ({len(lines)} lines) ---")
        add("\n".join(lines[:4] + (["..."] + lines[-2:] if len(lines) > 6 else lines[4:])))
        add(f"parsed timestamps: {len(read_latency(adb, cand))}")
    add("\n--- timestats ---")
    add(adb.sh("dumpsys SurfaceFlinger --timestats -enable").strip() or "(enable: no output)")
    time.sleep(1.5)
    add(f"game layers: {timestats_frames(adb, pkg)}")
    add("\n--- service call SurfaceFlinger 1013 ---")
    a = adb.sh("service call SurfaceFlinger 1013").strip()
    time.sleep(1)
    b = adb.sh("service call SurfaceFlinger 1013").strip()
    add(f"{a}\n{b}")
    add("\n--- gfxinfo ---")
    add(f"frames: {gfxinfo_frames(adb, pkg)}")
    add("\n--- SurfaceFlinger header ---")
    add("\n".join(adb.sh("dumpsys SurfaceFlinger | head -n 25").splitlines()[:25]))
    return "\n".join(rep)


# ------------------------------------------------------------------- CPU / GPU / mem
def read_cpu_raw(adb, pid):
    out = adb.sh(f"head -n1 /proc/stat; cat /proc/{pid}/stat 2>/dev/null")
    total = idle = proc = None
    for line in out.splitlines():
        if line.startswith("cpu "):
            v = [int(x) for x in line.split()[1:]]
            total = sum(v[:8])
            idle = v[3] + (v[4] if len(v) > 4 else 0)
        elif re.match(r"^\d+ \(", line):
            rest = line[line.rfind(")") + 2:].split()
            if len(rest) > 12:
                proc = int(rest[11]) + int(rest[12])
    return total, idle, proc


def read_cpu_top(adb, pid):
    """Fallback when /proc/stat is not readable: game CPU from top, normalized to all cores."""
    out = adb.sh(f"top -b -n 1 -p {pid}")
    m = re.search(r"(\d+)%cpu", out)
    cap = int(m.group(1)) if m else None
    idx = None
    for line in out.splitlines():
        toks = line.split()
        if idx is None and "PID" in toks:
            for i, t in enumerate(toks):
                if "%CPU" in t:
                    idx = i + (1 if t.startswith("S[") else 0)
                    break
        elif idx is not None and toks and toks[0] == str(pid) and len(toks) > idx:
            try:
                val = float(toks[idx])
            except ValueError:
                return None
            return val / cap if cap else val / 100
    return None


class CpuMeter:
    def __init__(self, adb):
        self.adb, self.prev, self.use_top = adb, None, False

    def sample(self, pid):
        if self.use_top:
            return read_cpu_top(self.adb, pid), None
        total, idle, proc = read_cpu_raw(self.adb, pid)
        if total is None:
            self.use_top = True
            return read_cpu_top(self.adb, pid), None
        game = total_cpu = None
        if self.prev and proc is not None and self.prev[2] is not None:
            dt = total - self.prev[0]
            if dt > 0:
                game = max(0.0, min(1.0, (proc - self.prev[2]) / dt))
                total_cpu = max(0.0, min(1.0, 1 - (idle - self.prev[1]) / dt))
        self.prev = (total, idle, proc)
        return game, total_cpu


def parse_gpu(path, txt):
    nums = [float(x) for x in re.findall(r"\d+(?:\.\d+)?", txt)]
    if not nums:
        return None
    if path.endswith("gpubusy"):
        return (nums[0] / nums[1]) if len(nums) >= 2 and nums[1] > 0 else 0.0
    return max(0.0, min(1.0, nums[0] / 100))


def detect_gpu_path(adb):
    script = "; ".join(f'echo "@@{p}"; cat {p} 2>/dev/null' for p in GPU_PATHS)
    out = adb.sh(script)
    for block in out.split("@@")[1:]:
        path, _, body = block.partition("\n")
        path = path.strip()
        if body.strip() and parse_gpu(path, body) is not None:
            return path
    return None


def read_gpu(adb, path):
    return parse_gpu(path, adb.sh(f"cat {path} 2>/dev/null")) if path else None


def mem_available_mb(adb):
    m = re.search(r"MemAvailable:\s*(\d+)", adb.sh("cat /proc/meminfo"))
    return round(int(m.group(1)) / 1024) if m else None


def game_pss_mb(adb, pkg):
    out = adb.sh(f"dumpsys meminfo {pkg}", timeout=20)
    m = (re.search(r"TOTAL PSS:\s*(\d+)", out)
         or re.search(r"^\s*TOTAL:?\s+(\d+)", out, re.M))
    return round(int(m.group(1)) / 1024) if m else None


def device_temp(adb):
    """Phone body/CPU temperature from thermal HAL (SKIN preferred, then CPU max). None if unavailable."""
    out = adb.sh("dumpsys thermalservice", timeout=15)
    part = out.split("Current temperatures from HAL", 1)
    if len(part) < 2:
        return None
    part = part[1].split("Current cooling devices", 1)[0]
    vals = {}
    for m in re.finditer(r"mValue=(-?[\d.]+),\s*mType=(-?\d+)", part):
        v, t = float(m.group(1)), int(m.group(2))
        if 0 < v < 150:
            vals.setdefault(t, []).append(v)
    for t in (3, 0):          # 3 = SKIN, 0 = CPU
        if vals.get(t):
            return round(max(vals[t]), 1)
    return None


def battery_temp(adb):
    m = re.search(r"temperature:\s*(\d+)", adb.sh("dumpsys battery"))
    return int(m.group(1)) / 10 if m else None


# ------------------------------------------------------------------------ helpers
def mean(xs):
    xs = [x for x in xs if x is not None]
    return statistics.mean(xs) if xs else None


def rnd(x, n=2):
    return None if x is None else round(x, n)


def fmt_clock(sec):
    sec = int(sec)
    return f"{sec // 60:02d}:{sec % 60:02d}"


def fps_stats(ts):
    if len(ts) < 3:
        return {}
    dur = (ts[-1] - ts[0]) / 1e9
    ft = [(b - a) / 1e6 for a, b in zip(ts, ts[1:]) if b > a]
    buckets = {}
    for t in ts:
        k = int((t - ts[0]) // 1_000_000_000)
        buckets[k] = buckets.get(k, 0) + 1
    full = int(dur)  # drop the last, incomplete second
    per_sec = [buckets.get(k, 0) for k in range(full)]
    ft_sorted = sorted(ft)
    p99 = ft_sorted[min(len(ft_sorted) - 1, int(len(ft_sorted) * 0.99))]
    return {
        "frames": len(ts),
        "avg_fps": (len(ts) - 1) / dur if dur > 0 else None,
        "min_fps": min(per_sec) if per_sec else None,
        "max_fps": max(per_sec) if per_sec else None,
        "avg_frame_time_ms": statistics.mean(ft) if ft else None,
        "p99_frame_time_ms": p99,
        "one_percent_low_fps": 1000 / p99 if p99 else None,
        "frames_over_50ms": sum(1 for x in ft if x > 50),
        "per_second": per_sec,
    }


def first_after(rows, key, sec):
    for r in rows:
        if r["t_sec"] >= sec and r.get(key) is not None:
            return r[key]
    return None


def last_value(rows, key):
    for r in reversed(rows):
        if r.get(key) is not None:
            return r[key]
    return None


def copy_to_clipboard(text):
    try:
        if os.name == "nt":
            subprocess.run(["clip"], input=text.encode("utf-16"), check=True)
        elif sys.platform == "darwin":
            env = dict(os.environ, LANG="en_US.UTF-8")
            subprocess.run(["pbcopy"], input=text.encode("utf-8"), env=env, check=True)
        else:
            return False
        return True
    except Exception:
        return False


def keep_awake():
    """macOS: stop the Mac from sleeping during the test (sleep would drop the USB/adb link)."""
    if sys.platform == "darwin" and shutil.which("caffeinate"):
        try:
            return subprocess.Popen(["caffeinate", "-dims", "-w", str(os.getpid())])
        except Exception:
            return None
    return None


def notify_done(title, text):
    if sys.platform == "darwin":
        subprocess.run(["afplay", "/System/Library/Sounds/Glass.aiff"], capture_output=True)
        script = f'display notification "{text}" with title "{title}"'
        subprocess.run(["osascript", "-e", script], capture_output=True)
    elif os.name == "nt":
        print("\a", end="", flush=True)


def open_folder(path):
    try:
        if sys.platform == "darwin":
            subprocess.run(["open", path], capture_output=True)
        elif os.name == "nt":
            os.startfile(path)
    except Exception:
        pass


# -------------------------------------------------------------------------- test
def run_test(adb, pkg, dev, args, out_root, log=None, on_sample=None, on_phase=None, stop_event=None):
    """Runs one game test. GUI passes log/on_sample/on_phase callbacks and a threading.Event to stop."""
    cli = log is None
    log = log or (lambda m: print(m))
    on_phase = on_phase or (lambda name, value=None: None)
    stop_event = stop_event or threading.Event()

    def wait(sec):
        return stop_event.wait(sec)  # True if stop requested

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_dir = os.path.join(out_root, f"{pkg}_{stamp}")
    os.makedirs(out_dir, exist_ok=True)
    events = []

    adb.sh("input keyevent KEYCODE_WAKEUP")
    component = launch_component(adb, pkg)
    if not component and not args.attach:
        log(f"  ! Cannot find launcher activity for {pkg}. Is the package name correct?")
        return None

    temp_before = battery_temp(adb)
    dtemp_before = device_temp(adb)
    t1 = t2 = None
    if args.attach:
        log("  Attach mode: launch times are skipped")
        avail_before = mem_available_mb(adb)
    else:
        adb.sh(f"am force-stop {pkg}")
        if wait(3):
            return None
        avail_before = mem_available_mb(adb)
        on_phase("launch1")
        log(f"  1st launch ({component}) ...")
        t1 = cold_start(adb, pkg, component)
        log(f"  1st launch: {t1} s")
        on_phase("launch1_done", t1)
        log(f"  waiting {args.between} s before 2nd launch ...")
        if wait(args.between):
            return None
        on_phase("launch2")
        log("  2nd launch ...")
        t2 = cold_start(adb, pkg, component)
        log(f"  2nd launch: {t2} s")
        on_phase("launch2_done", t2)

    pid = None
    for _ in range(20):
        pid = pidof(adb, pkg)
        if pid or wait(1):
            break
    if not pid:
        log("  ! Game process not found. Skipping.")
        return None

    gpu_path = detect_gpu_path(adb)
    log(f"  GPU load source: {gpu_path or 'not available on this device (column stays empty)'}")

    frames = FrameCollector(adb, pkg, log=log, force=getattr(args, "fps_method", None))
    frames.start()
    cpu = CpuMeter(adb)
    cpu.sample(pid)

    duration = args.duration * 60
    log(f">>> PLAY NOW. Test runs {args.duration:g} min." + (" Ctrl+C to stop earlier." if cli else ""))
    on_phase("play", duration)
    rows = []
    start = time.time()
    last_slow = -999
    pss = temp = dtemp = None
    missing = 0
    try:
        while not stop_event.is_set():
            now = time.time()
            el = now - start
            if el >= duration:
                break
            cur = pidof(adb, pkg)
            if not cur:
                missing += 1
                if missing >= 3:
                    events.append({"t_sec": round(el), "event": "game process disappeared (crash or closed)"})
                    log(f"  ! Game process disappeared at {fmt_clock(el)} - stopping")
                    break
            else:
                missing = 0
                if cur != pid:
                    events.append({"t_sec": round(el), "event": f"game restarted, new pid {cur}"})
                    pid, cpu.prev = cur, None
            game_cpu, total_cpu = cpu.sample(pid)
            gpu = read_gpu(adb, gpu_path)
            avail = mem_available_mb(adb)
            if el - last_slow >= 10:
                pss = game_pss_mb(adb, pkg)
                temp = battery_temp(adb)
                dtemp = device_temp(adb)
                last_slow = el
            fps = frames.fps_recent(max(1.0, args.interval))
            row = {
                "t_sec": round(el, 1), "fps": rnd(fps, 1), "cpu_game": rnd(game_cpu, 3),
                "cpu_total": rnd(total_cpu, 3), "gpu": rnd(gpu, 3), "mem_available_mb": avail,
                "game_pss_mb": pss, "battery_temp_c": temp, "device_temp_c": dtemp,
                "fps_method": frames.mode,
            }
            rows.append(row)
            if on_sample:
                on_sample(dict(row, duration=duration))
            if cli:
                print(f"\r  [{fmt_clock(el)}/{fmt_clock(duration)}] FPS {fps or 0:5.1f} | "
                      f"CPU game {(game_cpu or 0) * 100:4.0f}% | GPU "
                      f"{'n/a' if gpu is None else f'{gpu * 100:3.0f}%'} | RAM game {pss or '-'} MB | "
                      f"Batt {temp or '-'} C   ", end="", flush=True)
            wait(max(0.0, args.interval - (time.time() - now)))
        if stop_event.is_set():
            log("  Stopped by user")
            events.append({"t_sec": round(time.time() - start), "event": "stopped by user"})
    except KeyboardInterrupt:
        log("\n  Stopped by user")
        events.append({"t_sec": round(time.time() - start), "event": "stopped by user"})
    frames.stop_ev.set()
    frames.join(timeout=5)
    if cli:
        print()
    on_phase("saving")

    real_minutes = round((time.time() - start) / 60, 1)
    # one more temperature at the very end
    temp_end = battery_temp(adb)
    if rows:
        rows[-1]["battery_temp_c"] = temp_end
    fs = frames.stats()
    if not fs:
        log("  ! No FPS data. SurfaceFlinger did not return frames for this game.")

    gameplay_rows = [r for r in rows if r["t_sec"] >= 60] or rows
    avail_during = mean(r["mem_available_mb"] for r in gameplay_rows)
    temp_3 = first_after(rows, "battery_temp_c", 180)
    temp_30 = first_after(rows, "battery_temp_c", 1800)
    reached_30 = temp_30 is not None
    if not reached_30:
        temp_30 = temp_end
    # Y/Z: device skin temperature from the thermal HAL, battery temperature as a fallback
    dtemp_3 = first_after(rows, "device_temp_c", 180)
    if dtemp_3 is None:
        dtemp_3 = last_value(rows, "device_temp_c")
    y_temp = dtemp_before if dtemp_before is not None else temp_before
    z_temp = dtemp_3 if dtemp_before is not None else (temp_3 if temp_3 is not None else temp_end)
    # W: available RAM before minus during gameplay; falls back to game PSS when the difference is not positive
    pss_avg = rnd(mean(r["game_pss_mb"] for r in rows), 0)
    by_avail = (rnd(avail_before - avail_during, 0)
                if avail_before is not None and avail_during is not None else None)
    ram_used = by_avail if by_avail and by_avail > 0 else pss_avg

    summary = {
        "package": pkg, "component": component, "date": datetime.now().isoformat(timespec="seconds"),
        "tester": args.tester, "graphics": args.graphics, "device": dev,
        "duration_min": real_minutes,
        "launch_1_sec": t1, "launch_2_sec": t2,
        "avg_fps": rnd(fs.get("avg_fps"), 1), "min_fps": fs.get("min_fps"), "max_fps": fs.get("max_fps"),
        "avg_frame_time_ms": rnd(fs.get("avg_frame_time_ms"), 2),
        "one_percent_low_fps": rnd(fs.get("one_percent_low_fps"), 1),
        "frames_over_50ms": fs.get("frames_over_50ms"), "frames_total": fs.get("frames"),
        "fps_layer": frames.layer or frames.ts_layer, "fps_method": fs.get("method"),
        "avg_gpu": rnd(mean(r["gpu"] for r in rows), 2), "gpu_source": gpu_path,
        "avg_cpu_game": rnd(mean(r["cpu_game"] for r in rows), 2),
        "avg_cpu_total": rnd(mean(r["cpu_total"] for r in rows), 2),
        "battery_temp_before": temp_before, "battery_temp_after_30min": temp_30,
        "reached_30_min": reached_30, "battery_temp_after_3min": temp_3, "battery_temp_end": temp_end,
        "mem_available_before_mb": avail_before, "mem_available_during_mb": rnd(avail_during, 0),
        "game_ram_by_available_mb": by_avail, "game_ram_used_mb": ram_used,
        "game_ram_source": "available_before - available_during" if ram_used == by_avail else "game PSS",
        "game_pss_avg_mb": pss_avg,
        "device_temp_before": dtemp_before, "device_temp_after_3min": dtemp_3,
        "temp_before_test": y_temp, "temp_after_3min": z_temp,
        "game_pss_max_mb": max((r["game_pss_mb"] for r in rows if r["game_pss_mb"]), default=None),
        "events": events, "results_folder": os.path.abspath(out_dir),
    }

    row = [
        args.tester, args.graphics, dev["device"], dev["cpu"], dev["gpu"], dev["ram_gb"], dev["os"],
        f"{real_minutes:g} minute{'' if real_minutes == 1 else 's'}", t1, t2, summary["avg_fps"], summary["min_fps"], summary["max_fps"],
        summary["avg_frame_time_ms"], summary["avg_gpu"], summary["avg_cpu_game"], temp_before, temp_30,
        "", "", avail_before, summary["mem_available_during_mb"], ram_used,
        "", y_temp, z_temp,
    ]
    summary["table_row"] = dict(zip([c or f"(empty {i})" for i, c in enumerate(TABLE_COLUMNS)], row))

    save_outputs(out_dir, rows, fs, summary, row, args)
    return summary, row


def tsv_value(v, comma):
    if v is None:
        return ""
    s = f"{v:g}" if isinstance(v, float) else str(v)
    return s.replace(".", ",") if comma and isinstance(v, float) else s


def save_outputs(out_dir, rows, fs, summary, row, args):
    with open(os.path.join(out_dir, "samples.csv"), "w", newline="", encoding="utf-8") as f:
        if rows:
            w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
            w.writeheader()
            w.writerows(rows)
    with open(os.path.join(out_dir, "fps_per_second.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["second", "fps"])
        for i, v in enumerate(fs.get("per_second", [])):
            w.writerow([i, v])
    with open(os.path.join(out_dir, "summary.json"), "w", encoding="utf-8") as f:
        json.dump(summary, f, ensure_ascii=False, indent=2)
    line = "\t".join(tsv_value(v, args.decimal_comma) for v in row)
    with open(os.path.join(out_dir, "table_row.tsv"), "w", encoding="utf-8") as f:
        f.write("\t".join(TABLE_COLUMNS) + "\n" + line + "\n")
    summary["_tsv_line"] = line

    try:
        import openpyxl
        from openpyxl.styles import Font
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Performance results"
        ws.append(TABLE_COLUMNS)
        ws.append(row)
        for c in ws[1]:
            c.font = Font(bold=True)
        ws2 = wb.create_sheet("Samples")
        if rows:
            ws2.append(list(rows[0].keys()))
            for r in rows:
                ws2.append(list(r.values()))
        wb.save(os.path.join(out_dir, "summary.xlsx"))
    except ImportError:
        pass

    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        fig, axes = plt.subplots(3, 1, figsize=(11, 8), sharex=False)
        ps = fs.get("per_second", [])
        if ps:
            axes[0].plot(range(len(ps)), ps, lw=1)
        else:
            axes[0].text(0.5, 0.5, "No FPS data", ha="center", va="center", transform=axes[0].transAxes)
            axes[0].set_xticks([])
            axes[0].set_yticks([])
        axes[0].set_title(f"{summary['package']} - FPS per second (avg {summary['avg_fps']})")
        axes[0].set_ylabel("FPS")
        t = [r["t_sec"] for r in rows]
        axes[1].plot(t, [(r["cpu_game"] or 0) * 100 for r in rows], label="CPU game %")
        if any(r["gpu"] is not None for r in rows):
            axes[1].plot(t, [(r["gpu"] or 0) * 100 for r in rows], label="GPU %")
        axes[1].set_ylabel("%")
        axes[1].legend()
        axes[2].plot(t, [r["game_pss_mb"] for r in rows], label="Game RAM (PSS), MB")
        ax3 = axes[2].twinx()
        ax3.plot(t, [r["battery_temp_c"] for r in rows], color="tab:red", label="Battery C")
        axes[2].set_xlabel("seconds")
        axes[2].legend(loc="upper left")
        ax3.legend(loc="upper right")
        fig.tight_layout()
        fig.savefig(os.path.join(out_dir, "chart.png"), dpi=110)
        plt.close(fig)
    except ImportError:
        pass


# --------------------------------------------------------------------------- CLI
def pick_packages(adb, flt):
    out = adb.sh("pm list packages -3")
    pkgs = sorted(l.split(":", 1)[1].strip() for l in out.splitlines() if l.startswith("package:"))
    if flt:
        pkgs = [p for p in pkgs if flt.lower() in p.lower()]
    if not pkgs:
        sys.exit("No installed third-party packages found.")
    for i, p in enumerate(pkgs, 1):
        print(f"  {i:3d}. {p}")
    raw = input("\nGame numbers to test (e.g. 1,4,7): ")
    chosen = []
    for part in re.split(r"[,\s]+", raw.strip()):
        if part.isdigit() and 1 <= int(part) <= len(pkgs):
            chosen.append(pkgs[int(part) - 1])
    return chosen


def main():
    ap = argparse.ArgumentParser(description="Automatic Android game performance test via adb")
    ap.add_argument("packages", nargs="*", help="game package names (empty = choose from list)")
    ap.add_argument("-d", "--duration", type=float, default=30, help="test length in minutes (default 30)")
    ap.add_argument("-i", "--interval", type=float, default=2, help="sample interval, sec (default 2)")
    ap.add_argument("-t", "--tester", default="", help="tester name for the results row")
    ap.add_argument("-g", "--graphics", default="", help="graphic settings, e.g. Low/Normal/High")
    ap.add_argument("-s", "--serial", help="device serial if several are connected")
    ap.add_argument("--adb", default="adb", help="path to adb (default: adb from PATH)")
    ap.add_argument("--filter", help="filter for the package list, e.g. 'game'")
    ap.add_argument("--between", type=int, default=10, help="seconds between 1st and 2nd launch")
    ap.add_argument("--attach", action="store_true", help="game already running, skip launch timing")
    ap.add_argument("-y", "--yes", action="store_true", help="do not wait for Enter between games")
    ap.add_argument("--decimal-comma", action="store_true", help="use 7,94 instead of 7.94 in the TSV row")
    ap.add_argument("--fps-method", choices=list(FPS_METHOD_NAMES), help="force one FPS method")
    ap.add_argument("--diag", action="store_true", help="print FPS diagnostics for the (running) game and exit")
    ap.add_argument("-o", "--out", default=os.path.join(os.path.dirname(os.path.abspath(__file__)), "perf_results"),
                    help="output folder (default: perf_results next to the script)")
    args = ap.parse_args()

    args.adb = find_adb(args.adb)
    keep_awake()
    devs = list_devices(args.adb)
    if not devs:
        sys.exit("No device. Turn on USB debugging, connect the phone, accept the RSA prompt on the phone\n"
                 "and check with:  adb devices   (status must be 'device', not 'unauthorized')")
    serial = args.serial
    if not serial:
        if len(devs) > 1:
            for i, d in enumerate(devs, 1):
                print(f"  {i}. {d}")
            serial = devs[int(input("Device number: ")) - 1]
        else:
            serial = devs[0]
    adb = Adb(args.adb, serial)
    dev = device_info(adb)
    print(f"\nDevice: {dev['device']} | {dev['os']} | CPU {dev['cpu']} | GPU {dev['gpu']} | "
          f"RAM {dev['ram_gb']} GB\n")

    pkgs = args.packages or pick_packages(adb, args.filter)
    if args.diag:
        for pkg in pkgs:
            print(fps_diagnostics(adb, pkg))
        return
    if not pkgs:
        sys.exit("Nothing selected.")
    if not args.graphics:
        args.graphics = input("Graphic settings in the game (Low/Normal/High/Ultra): ").strip()

    session = os.path.join(args.out, datetime.now().strftime("session_%Y%m%d_%H%M%S"))
    os.makedirs(session, exist_ok=True)
    lines = []
    for n, pkg in enumerate(pkgs, 1):
        print(f"\n=== [{n}/{len(pkgs)}] {pkg} ===")
        if not args.yes:
            input("  Unlock the phone, then press Enter to start ...")
        res = run_test(adb, pkg, dev, args, session)
        if not res:
            continue
        summary, row = res
        lines.append(summary["_tsv_line"])
        print(f"  Avg FPS {summary['avg_fps']} | min {summary['min_fps']} | max {summary['max_fps']} | "
              f"1st launch {summary['launch_1_sec']} s | 2nd {summary['launch_2_sec']} s")
        print(f"  Results: {summary['results_folder']}")
        adb.sh(f"am force-stop {pkg}")
        notify_done("perfgrab", f"{pkg}: done, avg FPS {summary['avg_fps']}")

    if lines:
        all_path = os.path.join(session, "all_rows.tsv")
        with open(all_path, "w", encoding="utf-8") as f:
            f.write("\t".join(TABLE_COLUMNS) + "\n" + "\n".join(lines) + "\n")
        text = "\n".join(lines)
        copied = copy_to_clipboard(text)
        print("\nRow(s) for the 'Performance results' sheet (columns A-Z):\n")
        print(text)
        print(f"\nSaved: {os.path.abspath(all_path)}")
        if copied:
            key = "Cmd+V" if sys.platform == "darwin" else "Ctrl+V"
            print(f"Copied to clipboard - click cell A of an empty row in Google Sheets and press {key}.")
        open_folder(os.path.abspath(session))


if __name__ == "__main__":
    main()
